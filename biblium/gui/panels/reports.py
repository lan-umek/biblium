# -*- coding: utf-8 -*-
"""
Report Panels (obnovljeno v 2.19)
=================================
``biblium.gui.panels.reports`` je bil referenciran v ``gui/panels/__init__.py``
in ``gui/app.py`` (2.16, 2.18), a datoteka v mapah ni obstajala, zato se GUI
paket ni mogel uvoziti. Ta modul ga obnovi:

- ``CustomReportPanel``  — zbirka tabel/slik, ki jih uporabnik doda prek
  kontekstnega menija (``report_queue``): urejanje vrstnega reda, odstranitev,
  izvoz v HTML (samostojna datoteka), Word (python-docx) in Excel.
- ``ReportBuilderPanel`` — standardna poročila iz ``BiblioStats``: Excel z
  vsemi tabelami, HTML nadzorna plošča, PDF (če je na voljo), Methods text.
"""
from __future__ import annotations

import base64
import html as _html
import io
import os
import threading
import tkinter as tk
from tkinter import ttk, messagebox, filedialog
from typing import Any, Dict, List

from biblium.gui.config import FONTS
from biblium.gui.panels.base import BasePanel
from biblium.gui.widgets.cards import Card
from biblium.gui.widgets.buttons import ThemedButton, ActionButton
from biblium.gui.widgets.forms import LabeledEntry, LabeledCheckbox, LabeledCombobox
from biblium.gui.widgets.tables import DataTable

try:
    import pandas as pd
    HAS_PANDAS = True
except ImportError:  # pragma: no cover
    HAS_PANDAS = False


# ---------------------------------------------------------------------------
# Generiranje poročil iz čakalne vrste (neodvisno od Tk, testabilno)
# ---------------------------------------------------------------------------

def _item_rows(item) -> int:
    try:
        return len(item.data)
    except Exception:
        return 0


def build_html_report(items: List[Any], title: str = "Biblium report", subtitle: str = "",
                      max_rows: int = 200) -> str:
    """Sestavi samostojen HTML iz ``ReportItem``-ov (tabele kot HTML, slike kot base64 PNG)."""
    parts = [
        "<!doctype html><html><head><meta charset='utf-8'>",
        f"<title>{_html.escape(title)}</title>",
        "<style>body{font-family:Segoe UI,Arial,sans-serif;margin:32px;color:#1e293b}"
        "h1{margin-bottom:0}h2{margin-top:36px;border-bottom:1px solid #e2e8f0;padding-bottom:4px}"
        ".muted{color:#64748b}table{border-collapse:collapse;font-size:13px;margin:8px 0}"
        "th,td{border:1px solid #e2e8f0;padding:4px 8px;text-align:left}th{background:#f8fafc}"
        "img{max-width:100%;height:auto;border:1px solid #e2e8f0}nav li{margin:2px 0}</style></head><body>",
        f"<h1>{_html.escape(title)}</h1>",
    ]
    if subtitle:
        parts.append(f"<p class='muted'>{_html.escape(subtitle)}</p>")
    parts.append("<nav><ol>")
    for i, it in enumerate(items):
        parts.append(f"<li><a href='#item{i}'>{_html.escape(it.title)}</a> "
                     f"<span class='muted'>({it.item_type}, {_html.escape(it.source_panel)})</span></li>")
    parts.append("</ol></nav>")
    for i, it in enumerate(items):
        parts.append(f"<h2 id='item{i}'>{_html.escape(it.title)}</h2>")
        parts.append(f"<p class='muted'>Source: {_html.escape(it.source_panel)} · {it.timestamp}</p>")
        if it.description:
            parts.append(f"<p>{_html.escape(it.description)}</p>")
        if it.item_type == "table" and HAS_PANDAS and hasattr(it.data, "to_html"):
            df = it.data.head(max_rows)
            parts.append(df.to_html(index=False, border=0, na_rep=""))
            if len(it.data) > max_rows:
                parts.append(f"<p class='muted'>… {len(it.data) - max_rows} more rows (see Excel export)</p>")
        elif it.item_type == "plot":
            b = it.data
            if hasattr(b, "savefig"):
                buf = io.BytesIO(); b.savefig(buf, format="png", dpi=150, bbox_inches="tight"); b = buf.getvalue()
            if isinstance(b, (bytes, bytearray)):
                parts.append(f"<img src='data:image/png;base64,{base64.b64encode(b).decode('ascii')}'>")
    parts.append("</body></html>")
    return "\n".join(parts)


def save_report_items(items: List[Any], path: str, fmt: str = "html", title: str = "Biblium report",
                      subtitle: str = "") -> str:
    """Shrani čakalno vrsto kot html / docx / xlsx. Vrne pot."""
    fmt = fmt.lower()
    if fmt == "html":
        with open(path, "w", encoding="utf-8") as f:
            f.write(build_html_report(items, title=title, subtitle=subtitle))
    elif fmt == "xlsx":
        tables = [it for it in items if it.item_type == "table"]
        if not tables:
            raise ValueError("No tables in the report queue.")
        with pd.ExcelWriter(path) as xw:
            used = set()
            for i, it in enumerate(tables):
                name = "".join(ch for ch in it.title if ch.isalnum() or ch in " _-")[:28].strip() or f"table{i}"
                base, k = name, 1
                while name in used:
                    k += 1; name = f"{base[:25]}_{k}"
                used.add(name)
                it.data.to_excel(xw, sheet_name=name, index=False)
    elif fmt == "docx":
        from docx import Document
        from docx.shared import Inches
        doc = Document()
        doc.add_heading(title, 0)
        if subtitle:
            doc.add_paragraph(subtitle)
        for it in items:
            doc.add_heading(it.title, 1)
            doc.add_paragraph(f"Source: {it.source_panel} · {it.timestamp}")
            if it.description:
                doc.add_paragraph(it.description)
            if it.item_type == "table" and hasattr(it.data, "columns"):
                df = it.data.head(100)
                t = doc.add_table(rows=1, cols=len(df.columns))
                t.style = "Light Grid Accent 1" if "Light Grid Accent 1" in [s.name for s in doc.styles] else t.style
                for j, c in enumerate(df.columns):
                    t.rows[0].cells[j].text = str(c)
                for _, row in df.iterrows():
                    cells = t.add_row().cells
                    for j, v in enumerate(row):
                        cells[j].text = "" if (v is None or (isinstance(v, float) and v != v)) else str(v)
            elif it.item_type == "plot":
                b = it.data
                if hasattr(b, "savefig"):
                    buf = io.BytesIO(); b.savefig(buf, format="png", dpi=150, bbox_inches="tight"); b = buf.getvalue()
                if isinstance(b, (bytes, bytearray)):
                    doc.add_picture(io.BytesIO(bytes(b)), width=Inches(6))
        doc.save(path)
    else:
        raise ValueError(f"Unknown format: {fmt}")
    return path


# ---------------------------------------------------------------------------
# Custom report (čakalna vrsta)
# ---------------------------------------------------------------------------

class CustomReportPanel(BasePanel):
    """Report builder from queued tables and plots."""

    title = "Custom Report Builder"
    icon = "📝"
    description = "Assemble a report from tables and plots added via right-click → 'Add to report'."
    requires_data = False

    def __init__(self, parent, theme: str = "light", bib=None, **kwargs):
        super().__init__(parent, theme=theme, bib=bib, **kwargs)
        self._primary_action = self._refresh_queue

    def _create_options(self):
        self._add_title()
        card = Card(self.options_content, title="📄 Report", theme=self.theme_name)
        card.pack(fill=tk.X, padx=8, pady=8)
        self.title_entry = LabeledEntry(card.content, label="Title:", default="Bibliometric report",
                                        theme=self.theme_name, label_width=10)
        self.title_entry.pack(fill=tk.X, pady=3)
        self.subtitle_entry = LabeledEntry(card.content, label="Subtitle:", default="",
                                           theme=self.theme_name, label_width=10)
        self.subtitle_entry.pack(fill=tk.X, pady=3)
        self.format_combo = LabeledCombobox(card.content, label="Format:", values=["html", "docx", "xlsx"],
                                            default="html", theme=self.theme_name, label_width=10)
        self.format_combo.pack(fill=tk.X, pady=3)

        btns = tk.Frame(self.options_content, bg=self.theme["bg_secondary"])
        btns.pack(fill=tk.X, padx=8, pady=12)
        ActionButton(btns, text="Generate Report", icon="📥", command=self._generate,
                     theme=self.theme_name).pack(fill=tk.X)
        ThemedButton(btns, text="Refresh Queue", icon="🔄", style="secondary", command=self._refresh_queue,
                     theme=self.theme_name).pack(fill=tk.X, pady=(8, 0))
        row = tk.Frame(btns, bg=self.theme["bg_secondary"])
        row.pack(fill=tk.X, pady=(8, 0))
        for text, cmd in (("▲ Up", self._move_up), ("▼ Down", self._move_down), ("✖ Remove", self._remove)):
            ThemedButton(row, text=text, style="secondary", command=cmd, theme=self.theme_name).pack(
                side=tk.LEFT, expand=True, fill=tk.X, padx=2)
        ThemedButton(btns, text="Clear Queue", style="secondary", command=self._clear,
                     theme=self.theme_name).pack(fill=tk.X, pady=(8, 0))

    def _create_results(self):
        self.results_content = tk.Frame(self.results_frame, bg=self.theme["bg_card"])
        self.results_content.pack(fill=tk.BOTH, expand=True, padx=8, pady=8)
        self.table = DataTable(self.results_content, theme=self.theme_name, show_ai_button=False)
        self.table.pack(fill=tk.BOTH, expand=True)
        self._refresh_queue()

    def _queue(self):
        from biblium.gui.core.state import report_queue
        return report_queue

    def _refresh_queue(self):
        items = self._queue().get_items()
        rows = [{"#": i + 1, "Type": it.item_type, "Title": it.title, "Source": it.source_panel,
                 "Rows": _item_rows(it) if it.item_type == "table" else "", "Added": it.timestamp}
                for i, it in enumerate(items)]
        if HAS_PANDAS:
            self.table.set_data(pd.DataFrame(rows, columns=["#", "Type", "Title", "Source", "Rows", "Added"]))

    def _selected_index(self):
        try:
            sel = self.table.get_selected_rows()
            if sel:
                return int(sel[0]["#"]) - 1
        except Exception:
            pass
        return None

    def _move_up(self):
        i = self._selected_index()
        if i is not None:
            self._queue().move_up(i); self._refresh_queue()

    def _move_down(self):
        i = self._selected_index()
        if i is not None:
            self._queue().move_down(i); self._refresh_queue()

    def _remove(self):
        i = self._selected_index()
        if i is not None:
            self._queue().remove_by_index(i); self._refresh_queue()

    def _clear(self):
        if messagebox.askyesno("Clear", "Remove all items from the report queue?"):
            self._queue().clear(); self._refresh_queue()

    def _generate(self):
        items = self._queue().get_items()
        if not items:
            messagebox.showinfo("Report", "The report queue is empty. Right-click a table or plot → 'Add to report'.")
            return
        fmt = self.format_combo.get() or "html"
        path = filedialog.asksaveasfilename(defaultextension=f".{fmt}",
                                            filetypes=[(fmt.upper(), f"*.{fmt}")],
                                            initialfile=f"biblium_report.{fmt}")
        if not path:
            return
        try:
            save_report_items(items, path, fmt=fmt, title=self.title_entry.get() or "Biblium report",
                              subtitle=self.subtitle_entry.get() or "")
            messagebox.showinfo("Report", f"Report saved to:\n{path}")
        except Exception as e:  # noqa: BLE001
            messagebox.showerror("Report", str(e))


# ---------------------------------------------------------------------------
# Standard report builder
# ---------------------------------------------------------------------------

class ReportBuilderPanel(BasePanel):
    """Standard reports generated from the loaded BiblioStats object."""

    title = "Report Builder"
    icon = "📑"
    description = "Generate standard outputs from the loaded dataset (Excel tables, HTML dashboard, PDF, methods text)."
    requires_data = True

    def __init__(self, parent, theme: str = "light", bib=None, **kwargs):
        super().__init__(parent, theme=theme, bib=bib, **kwargs)
        self._primary_action = self._generate

    def _create_options(self):
        self._add_title()
        card = Card(self.options_content, title="📦 Outputs", theme=self.theme_name)
        card.pack(fill=tk.X, padx=8, pady=8)
        self.cb_excel = LabeledCheckbox(card.content, label="Excel workbook with all computed tables",
                                        default=True, theme=self.theme_name)
        self.cb_excel.pack(fill=tk.X, pady=2)
        self.cb_dashboard = LabeledCheckbox(card.content, label="HTML dashboard", default=True,
                                            theme=self.theme_name)
        self.cb_dashboard.pack(fill=tk.X, pady=2)
        self.cb_pdf = LabeledCheckbox(card.content, label="PDF report", default=False, theme=self.theme_name)
        self.cb_pdf.pack(fill=tk.X, pady=2)
        self.cb_methods = LabeledCheckbox(card.content, label="Methods paragraph (txt, en + sl)", default=True,
                                          theme=self.theme_name)
        self.cb_methods.pack(fill=tk.X, pady=2)
        self.title_entry = LabeledEntry(card.content, label="Title:", default="Bibliometric Analysis Report",
                                        theme=self.theme_name, label_width=8)
        self.title_entry.pack(fill=tk.X, pady=3)
        btns = tk.Frame(self.options_content, bg=self.theme["bg_secondary"])
        btns.pack(fill=tk.X, padx=8, pady=12)
        self.run_btn = ActionButton(btns, text="Generate…", icon="▶", command=self._generate, theme=self.theme_name)
        self.run_btn.pack(fill=tk.X)

    def _create_results(self):
        self.results_content = tk.Frame(self.results_frame, bg=self.theme["bg_card"])
        self.results_content.pack(fill=tk.BOTH, expand=True, padx=8, pady=8)
        self.log = tk.Text(self.results_content, font=FONTS.get_font("monospace"), bg=self.theme["bg_card"],
                           fg=self.theme["text_primary"], relief=tk.FLAT, padx=12, pady=12, wrap=tk.WORD)
        self.log.pack(fill=tk.BOTH, expand=True)
        self._log("Choose outputs and click Generate…")

    def _log(self, msg):
        try:
            self.log.config(state=tk.NORMAL); self.log.insert(tk.END, msg + "\n"); self.log.config(state=tk.DISABLED)
        except Exception:
            pass

    def _generate(self):
        if self.bib is None:
            messagebox.showwarning("No Data", "Please load a dataset first.")
            return
        folder = filedialog.askdirectory(title="Output folder")
        if not folder:
            return
        want = {"excel": self.cb_excel.get(), "dashboard": self.cb_dashboard.get(),
                "pdf": self.cb_pdf.get(), "methods": self.cb_methods.get()}
        title = self.title_entry.get() or "Bibliometric Analysis Report"
        self.run_btn.config(state=tk.DISABLED, text="⏳ Generating...")

        def work():
            outputs: List[str] = []
            errors: List[str] = []
            try:
                if want["excel"]:
                    p = os.path.join(folder, "biblium_tables.xlsx")
                    try:
                        self.bib.export_all_tables(p); outputs.append(p)
                    except Exception as e:  # noqa: BLE001
                        errors.append(f"Excel: {e}")
                if want["dashboard"]:
                    p = os.path.join(folder, "dashboard.html")
                    try:
                        self.bib.create_dashboard(p, title=title); outputs.append(p)
                    except Exception as e:  # noqa: BLE001
                        errors.append(f"Dashboard: {e}")
                if want["pdf"]:
                    p = os.path.join(folder, "bibliometric_report.pdf")
                    try:
                        self.bib.export_pdf_report(p, title=title); outputs.append(p)
                    except Exception as e:  # noqa: BLE001
                        errors.append(f"PDF: {e}")
                if want["methods"]:
                    p = os.path.join(folder, "methods.txt")
                    try:
                        with open(p, "w", encoding="utf-8") as f:
                            f.write(self.bib.methods_text("en") + "\n\n" + self.bib.methods_text("sl") + "\n")
                        outputs.append(p)
                    except Exception as e:  # noqa: BLE001
                        errors.append(f"Methods: {e}")
            finally:
                self._safe_after(lambda: self._done(outputs, errors))

        threading.Thread(target=work, daemon=True).start()

    def _done(self, outputs, errors):
        self.run_btn.config(state=tk.NORMAL, text="▶ Generate…")
        for p in outputs:
            self._log(f"✓ {p}")
        for e in errors:
            self._log(f"✗ {e}")
        if outputs and not errors:
            messagebox.showinfo("Report", f"Generated {len(outputs)} file(s).")
        elif errors:
            messagebox.showwarning("Report", "\n".join(errors))


__all__ = ["CustomReportPanel", "ReportBuilderPanel", "build_html_report", "save_report_items"]
